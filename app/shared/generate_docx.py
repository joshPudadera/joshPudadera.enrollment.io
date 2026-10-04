"""
generate_docx.py
================
Builds a Word admission summary document from all data the student
submitted in the online enrollment form (stored in pre_registrations).

Called by generate_docx.php:
    python generate_docx.py <payload_json_path> <output_docx_path>

Requires: pip install python-docx
"""

import sys
import json
import os

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("ERROR: python-docx not installed. Run: pip install python-docx", file=sys.stderr)
    sys.exit(1)


# ── Colours ───────────────────────────────────────────────────
BLUE  = RGBColor(0x1a, 0x3a, 0x8c)
DARK  = RGBColor(0x1a, 0x1a, 0x2e)
GREY  = RGBColor(0x55, 0x55, 0x55)
LIGHT = RGBColor(0x88, 0x88, 0x88)


def v(data: dict, *keys) -> str:
    """Safely get a nested value, return em-dash if missing/empty."""
    for k in keys:
        if isinstance(data, dict):
            data = data.get(k)
        else:
            return '—'
    return str(data).strip() if data else '—'


def hr(doc):
    p    = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),   'single')
    bot.set(qn('w:sz'),    '4')
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), '1a3a8c')
    pBdr.append(bot)
    pPr.append(pBdr)


def section_heading(doc, title: str):
    hr(doc)
    p   = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(6)
    r   = p.add_run(title.upper())
    r.bold           = True
    r.font.size      = Pt(9.5)
    r.font.color.rgb = BLUE


def field_table(doc, rows: list):
    """Add a two-column label/value table."""
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style     = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    for label, value in rows:
        row = tbl.add_row()
        # Label cell
        lc = row.cells[0]
        lc.width = Cm(5)
        lp = lc.paragraphs[0]
        lp.paragraph_format.space_before = Pt(2)
        lp.paragraph_format.space_after  = Pt(2)
        lr = lp.add_run(label)
        lr.bold           = True
        lr.font.size      = Pt(8.5)
        lr.font.color.rgb = GREY
        # Value cell
        vc = row.cells[1]
        vp = vc.paragraphs[0]
        vp.paragraph_format.space_before = Pt(2)
        vp.paragraph_format.space_after  = Pt(2)
        vr = vp.add_run(value)
        vr.font.size      = Pt(9)
        vr.font.color.rgb = DARK if value != '—' else LIGHT


def build(payload: dict, output_path: str):
    doc = Document()

    # ── Page margins ─────────────────────────────────────────
    for sec in doc.sections:
        sec.top_margin    = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin   = Cm(2.5)
        sec.right_margin  = Cm(2.5)

    a   = payload.get('applicant', {})
    meta = payload

    # ── Header ───────────────────────────────────────────────
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(2)
    tr = t.add_run('BESTLINK COLLEGE OF THE PHILIPPINES')
    tr.bold = True; tr.font.size = Pt(14); tr.font.color.rgb = BLUE

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_after = Pt(2)
    sr = s.add_run('Office of the Registrar — Student Admission Summary')
    sr.font.size = Pt(10); sr.font.color.rgb = BLUE

    m = doc.add_paragraph()
    m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    m.paragraph_format.space_after = Pt(10)
    mr = m.add_run(
        f"Generated: {meta.get('generated_at','')}   |   By: {meta.get('generated_by','')}"
    )
    mr.font.size = Pt(8); mr.font.color.rgb = LIGHT

    # ── 1. Applicant Identity ─────────────────────────────────
    section_heading(doc, '1. Applicant Identity')
    full = ' '.join(filter(None, [v(a,'last_name'), v(a,'first_name'), v(a,'middle_name')]))
    if v(a,'suffix') != '—':
        full += ' ' + v(a,'suffix')
    field_table(doc, [
        ('Reference No.',       v(a, 'ref_number')),
        ('Full Name',           full if full.strip() and full.strip() != '—' else '—'),
        ('Last Name',           v(a, 'last_name')),
        ('First Name',          v(a, 'first_name')),
        ('Middle Name',         v(a, 'middle_name')),
        ('Suffix',              v(a, 'suffix')),
        ('Applicant Type',      v(a, 'applicant_type')),
        ('Application Status',  v(a, 'status')),
        ('Date Submitted',      v(a, 'submitted_at')),
    ])

    # ── 2. Personal Details ───────────────────────────────────
    section_heading(doc, '2. Personal Details')
    field_table(doc, [
        ('Date of Birth',   v(a, 'birthday')),
        ('Sex',             v(a, 'sex')),
        ('Civil Status',    v(a, 'civil_status')),
        ('Nationality',     v(a, 'nationality')),
        ('Religion',        v(a, 'religion')),
        ('Place of Birth',  v(a, 'place_of_birth')),
    ])

    # ── 3. Contact Information ────────────────────────────────
    section_heading(doc, '3. Contact Information')
    field_table(doc, [
        ('Email Address',   v(a, 'email')),
        ('Mobile Number',   v(a, 'phone')),
        ('Home Address',    v(a, 'address')),
    ])

    # ── 4. Academic Information ───────────────────────────────
    section_heading(doc, '4. Academic Information')
    field_table(doc, [
        ('Campus / Branch',      v(a, 'branch')),
        ('Course',               v(a, 'course')),
        ('Year Level',           v(a, 'year_level')),
        ('Transfer Year Level',  v(a, 'transfer_year_level')),
        ('Previous School',      v(a, 'prev_school')),
        ('School Year Graduated',v(a, 'grad_year')),
    ])

    # ── 5. Emergency Contact ──────────────────────────────────
    section_heading(doc, '5. Emergency Contact')
    field_table(doc, [
        ('Contact Person',  v(a, 'emergency_name')),
        ('Relationship',    v(a, 'emergency_relation')),
        ('Contact Number',  v(a, 'emergency_phone')),
    ])

    # ── 6. Submitted Documents ────────────────────────────────
    section_heading(doc, '6. Submitted Documents')
    docs_list = payload.get('documents', [])
    if docs_list:
        tbl = doc.add_table(rows=1, cols=4)
        tbl.style     = 'Table Grid'
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        hdr = tbl.rows[0].cells
        for i, h in enumerate(['Document Type', 'File Name', 'Status', 'Uploaded']):
            p = hdr[i].paragraphs[0]
            r = p.add_run(h)
            r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = BLUE
        for d in docs_list:
            row  = tbl.add_row()
            vals = [
                d.get('document_type', '—'),
                d.get('file_name', '—'),
                d.get('status', '—'),
                d.get('uploaded_at', '—')[:10] if d.get('uploaded_at') else '—',
            ]
            for i, val in enumerate(vals):
                p = row.cells[i].paragraphs[0]
                r = p.add_run(str(val))
                r.font.size = Pt(8.5)
    else:
        p = doc.add_paragraph()
        r = p.add_run('No documents uploaded yet.')
        r.font.size = Pt(9); r.font.color.rgb = LIGHT

    # ── Footer note ───────────────────────────────────────────
    doc.add_paragraph()
    fn = doc.add_paragraph()
    fn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fn.paragraph_format.space_before = Pt(14)
    fr = fn.add_run(
        'This document was system-generated by the BCP Enrollment Management System. '
        'All information was provided by the applicant during the online admission process.'
    )
    fr.font.size = Pt(7.5); fr.font.color.rgb = LIGHT; fr.italic = True

    doc.save(output_path)
    print(f'OK: {output_path}')


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print('Usage: python generate_docx.py <payload.json> <output.docx>', file=sys.stderr)
        sys.exit(1)

    json_path   = sys.argv[1]
    output_path = sys.argv[2]

    if not os.path.exists(json_path):
        print(f'Payload not found: {json_path}', file=sys.stderr)
        sys.exit(1)

    with open(json_path, 'r', encoding='utf-8-sig') as f:
        payload = json.load(f)

    build(payload, output_path)
